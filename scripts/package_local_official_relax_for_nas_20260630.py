from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_ROOT = Path(r"C:\Users\dawso\Desktop\LOCAL_DFTPY_OFFICIAL_RELAX_VACLM_10X10_20260629")
PACKAGE_ROOT = Path(r"C:\Users\dawso\Desktop\DFTPY_VACLM_OFFICIAL_RELAX_NAS_20260630")
REPO_ROOT = Path(r"C:\Users\dawso\nano_tensile_TFvW")

MU_VALUES = [round(i / 10, 1) for i in range(1, 11)]
LAM_VALUES = [round(i / 10, 1) for i in range(1, 11)]


def reset_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def copytree_contents(src: Path, dst: Path) -> None:
    dst.mkdir(parents=True, exist_ok=True)
    for item in src.iterdir():
        target = dst / item.name
        if item.is_dir():
            shutil.copytree(item, target, dirs_exist_ok=True)
        else:
            shutil.copy2(item, target)


def load_result(path: Path) -> dict | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def collect_rows() -> list[dict]:
    rows: list[dict] = []
    runs = SOURCE_ROOT / "03_runs"
    for case_dir in sorted(p for p in runs.iterdir() if p.is_dir() and p.name.startswith("tfvw_")):
        result = load_result(case_dir / "result.json")
        failed = (case_dir / "LOCAL_OFFICIAL_RERUN_FAILED.txt").exists()
        manifest = load_result(case_dir / "point_manifest.json") or {}
        lam = manifest.get("lambda")
        mu = manifest.get("mu")
        row = {
            "setting": case_dir.name,
            "lambda": lam,
            "mu": mu,
            "status": "ok" if result else ("failed" if failed else "missing"),
            "source_dir": str(case_dir),
            "result_json": str(case_dir / "result.json") if result else "",
            "pristine_traj": str(case_dir / "pristine_official_relax.traj") if (case_dir / "pristine_official_relax.traj").exists() else "",
            "vacancy_traj": str(case_dir / "vacancy_official_relax.traj") if (case_dir / "vacancy_official_relax.traj").exists() else "",
            "pristine_relax_log": str(case_dir / "pristine_official_relax.log") if (case_dir / "pristine_official_relax.log").exists() else "",
            "vacancy_relax_log": str(case_dir / "vacancy_official_relax.log") if (case_dir / "vacancy_official_relax.log").exists() else "",
        }
        if result:
            row.update(
                {
                    "total_formation_energy_eV": result.get("vacancy_formation_energy_eV"),
                    "kedf_formation_energy_eV": result.get("vacancy_formation_kedf_energy_eV"),
                    "lattice_constant_A": result.get("lattice_constant_A"),
                    "pristine_total_energy_eV": result.get("pristine_energy_eV"),
                    "vacancy_total_energy_eV": result.get("vacancy_energy_eV"),
                    "pristine_kedf_energy_eV": result.get("pristine_kedf_energy_eV"),
                    "vacancy_kedf_energy_eV": result.get("vacancy_kedf_energy_eV"),
                    "pristine_final_fmax_eV_A": (result.get("pristine") or {}).get("final_fmax_eV_A"),
                    "vacancy_final_fmax_eV_A": (result.get("vacancy") or {}).get("final_fmax_eV_A"),
                    "pristine_converged": (result.get("pristine") or {}).get("converged"),
                    "vacancy_converged": (result.get("vacancy") or {}).get("converged"),
                }
            )
        else:
            row.update(
                {
                    "total_formation_energy_eV": "",
                    "kedf_formation_energy_eV": "",
                    "lattice_constant_A": "",
                    "pristine_total_energy_eV": "",
                    "vacancy_total_energy_eV": "",
                    "pristine_kedf_energy_eV": "",
                    "vacancy_kedf_energy_eV": "",
                    "pristine_final_fmax_eV_A": "",
                    "vacancy_final_fmax_eV_A": "",
                    "pristine_converged": "",
                    "vacancy_converged": "",
                }
            )
        rows.append(row)
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0].keys()) if rows else []
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def matrix(rows: list[dict], key: str) -> list[list[object]]:
    by_pair = {}
    for row in rows:
        try:
            by_pair[(round(float(row["lambda"]), 1), round(float(row["mu"]), 1))] = row
        except Exception:
            continue
    table = [["lambda/mu", *MU_VALUES]]
    for lam in LAM_VALUES:
        line = [lam]
        for mu in MU_VALUES:
            row = by_pair.get((lam, mu))
            if not row:
                line.append("MISSING")
            elif row["status"] != "ok":
                line.append("FAILED")
            else:
                value = row.get(key)
                line.append(value if value not in (None, "") else "MISSING")
        table.append(line)
    return table


def write_matrix_csv(path: Path, table: list[list[object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(table)


def write_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Professor_Table"

    sections = [
        ("Vacancy formation energy (eV)", "total_formation_energy_eV"),
        ("KEDF formation contribution (eV)", "kedf_formation_energy_eV"),
        ("Relaxed pristine lattice constant (A)", "lattice_constant_A"),
    ]
    col = 1
    for title, key in sections:
        ws.cell(1, col, title)
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(1, col).fill = PatternFill("solid", fgColor="D9EAF7")
        table = matrix(rows, key)
        for r, line in enumerate(table, start=2):
            for c, value in enumerate(line, start=col):
                ws.cell(r, c, value)
                ws.cell(r, c).alignment = Alignment(horizontal="center")
        col += 12

    ws2 = wb.create_sheet("Flat_Results")
    fields = list(rows[0].keys()) if rows else []
    ws2.append(fields)
    for row in rows:
        ws2.append([row.get(field, "") for field in fields])

    ws3 = wb.create_sheet("Notes")
    notes = [
        ("workflow", "Official-style DFTpy relaxation rerun: DFTpyCalculator + ASE UnitCellFilter + BFGS."),
        ("official_reference", "https://dftpy.rutgers.edu/tutorials/ofdft/relax.html"),
        ("formation_energy_formula", "E_f^vac = E_vac(Al107) - (107/108) E_pristine(Al108)"),
        ("kedf_formula", "KEDF_f^vac = KEDF_vac(Al107) - (107/108) KEDF_pristine(Al108)"),
        ("lattice_constant", "Relaxed pristine 3x3x3 cell length divided by 3."),
        ("source", str(SOURCE_ROOT)),
    ]
    for row in notes:
        ws3.append(row)

    for sheet in wb.worksheets:
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 50)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_readmes(rows: list[dict]) -> None:
    ok = sum(1 for r in rows if r["status"] == "ok")
    failed = sum(1 for r in rows if r["status"] == "failed")
    traj_count = len(list((SOURCE_ROOT / "03_runs").rglob("*.traj")))
    result_count = len(list((SOURCE_ROOT / "03_runs").rglob("result.json")))

    (PACKAGE_ROOT / "README_PACKAGE.md").write_text(
        f"""# DFTpy VACLM Official-Style Relaxation Package

This package is a local verification rerun of the Al single-vacancy TFvW
lambda-mu scan using the DFTpy official relaxation style:

- `DFTpyCalculator(config=conf)`
- `calctype = Energy Force Stress`
- ASE `UnitCellFilter`
- ASE `BFGS`
- explicit ASE trajectory files

Official DFTpy reference:

```text
https://dftpy.rutgers.edu/tutorials/ofdft/relax.html
```

## Status

- attempted points: 100
- completed result.json: {result_count}
- successful cases: {ok}
- failed/pathological cases: {failed}
- trajectory files: {traj_count}

## Main files

- `01_RAW_TABLE/official_relax_professor_table.xlsx`
- `01_RAW_TABLE/official_relax_flat_results.csv`
- `03_RAW_CASES/03_runs/*`
- `04_RUN_METADATA/local_official_rerun_progress.json`
- `04_RUN_METADATA/local_official_rerun_summary.csv`
- `05_SCRIPTS_USED/run_dftpy_official_relax_vaclm_one.py`
- `05_SCRIPTS_USED/run_local_dftpy_official_relax_vaclm_100.py`

## Recommendation

Use this package as an official-style local verification dataset because it
contains `.traj` files and follows the DFTpy relaxation tutorial structure
more directly than the earlier iService delivery package.
""",
        encoding="utf-8",
    )

    (PACKAGE_ROOT / "04_RUN_METADATA" / "LOCAL_VS_ISERVICE_RECOMMENDATION.md").write_text(
        """# Local Official Rerun vs iService Package

## iService package

Strengths:

- HPC production provenance.
- Original professor-requested raw table already exists.
- Remote workflow is vc-relax equivalent: DFTpyCalculator + FrechetCellFilter + ASE optimizer.

Weaknesses:

- The local copied delivery package did not contain `.traj` files.
- The `.ini` files alone look like Optdensity inputs unless the Python/ASE runner is inspected.

## Local official-style rerun

Strengths:

- Follows the DFTpy relaxation tutorial style explicitly.
- Writes ASE `.traj` files for pristine and vacancy calculations.
- One case folder contains input, output, trajectory, final structures, logs, and result.json.
- Cleaner for advisor review of "was this relaxation?".

Weaknesses:

- It is a local verification rerun, not the original iService production run.
- Five pathological points failed under the guard and are marked explicitly.

## Recommendation

For the next advisor upload, use the local official-style rerun as the clean
verification package.  Keep the iService package as supporting HPC provenance,
especially if the remote `.traj` files can still be pulled.
""",
        encoding="utf-8",
    )


def main() -> None:
    if not SOURCE_ROOT.exists():
        raise FileNotFoundError(SOURCE_ROOT)
    rows = collect_rows()
    if len(rows) != 100:
        raise RuntimeError(f"Expected 100 cases, got {len(rows)}")

    reset_dir(PACKAGE_ROOT)
    for sub in [
        "01_RAW_TABLE",
        "02_SIMPLE_MAPS",
        "03_RAW_CASES",
        "04_RUN_METADATA",
        "05_SCRIPTS_USED",
        "06_REFERENCE_NOTES",
        "07_EVALUATION_SLIDES_PLACEHOLDER",
    ]:
        (PACKAGE_ROOT / sub).mkdir(parents=True, exist_ok=True)

    write_csv(PACKAGE_ROOT / "01_RAW_TABLE" / "official_relax_flat_results.csv", rows)
    write_matrix_csv(PACKAGE_ROOT / "01_RAW_TABLE" / "matrix_vacancy_formation_energy_eV.csv", matrix(rows, "total_formation_energy_eV"))
    write_matrix_csv(PACKAGE_ROOT / "01_RAW_TABLE" / "matrix_kedf_formation_energy_eV.csv", matrix(rows, "kedf_formation_energy_eV"))
    write_matrix_csv(PACKAGE_ROOT / "01_RAW_TABLE" / "matrix_lattice_constant_A.csv", matrix(rows, "lattice_constant_A"))
    write_xlsx(PACKAGE_ROOT / "01_RAW_TABLE" / "official_relax_professor_table.xlsx", rows)

    copytree_contents(SOURCE_ROOT / "03_runs", PACKAGE_ROOT / "03_RAW_CASES" / "03_runs")
    for name in [
        "BACKGROUND_PROCESS.json",
        "BACKGROUND_stdout.log",
        "BACKGROUND_stderr.log",
        "local_official_rerun_progress.json",
        "local_official_rerun_summary.csv",
        "RUN_DESCRIPTION.md",
        "official_relax_completed_95_summary.csv",
    ]:
        src = SOURCE_ROOT / name
        if src.exists():
            shutil.copy2(src, PACKAGE_ROOT / "04_RUN_METADATA" / name)

    for name in [
        "run_dftpy_official_relax_vaclm_one.py",
        "run_local_dftpy_official_relax_vaclm_100.py",
    ]:
        shutil.copy2(REPO_ROOT / "scripts" / name, PACKAGE_ROOT / "05_SCRIPTS_USED" / name)

    (PACKAGE_ROOT / "06_REFERENCE_NOTES" / "FORMULAS_AND_REFERENCES.md").write_text(
        """# Formulas and References

DFTpy relaxation reference:

```text
https://dftpy.rutgers.edu/tutorials/ofdft/relax.html
```

Vacancy formation energy:

```text
E_f^vac = E_vac(Al107) - (107/108) E_pristine(Al108)
```

KEDF formation contribution:

```text
KEDF_f^vac = KEDF_vac(Al107) - (107/108) KEDF_pristine(Al108)
```

Lattice constant:

```text
a0 = mean(relaxed pristine cell lengths) / 3
```
""",
        encoding="utf-8",
    )

    write_readmes(rows)

    audit = {
        "package_root": str(PACKAGE_ROOT),
        "source_root": str(SOURCE_ROOT),
        "case_dirs": len(list((PACKAGE_ROOT / "03_RAW_CASES" / "03_runs").iterdir())),
        "result_json": len(list((PACKAGE_ROOT / "03_RAW_CASES" / "03_runs").rglob("result.json"))),
        "failed_markers": len(list((PACKAGE_ROOT / "03_RAW_CASES" / "03_runs").rglob("LOCAL_OFFICIAL_RERUN_FAILED.txt"))),
        "traj_files": len(list((PACKAGE_ROOT / "03_RAW_CASES" / "03_runs").rglob("*.traj"))),
        "xlsx": str(PACKAGE_ROOT / "01_RAW_TABLE" / "official_relax_professor_table.xlsx"),
    }
    (PACKAGE_ROOT / "04_RUN_METADATA" / "PACKAGE_AUDIT.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
